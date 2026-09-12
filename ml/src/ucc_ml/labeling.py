"""Blind labelling for the UCC ML Lab (contracts K3, K13, K14).

Round-aware from the start: the pilot round (pilot_v1) and the main round (main_v1) run through the
same functions. The module grows over Tasks 13-19:
  Task 13  round paths, the chunked blind queue with hidden repeats, its private key, pre-registration
  Task 15  labeller briefs, structured-output validation, raw labeller CSVs, labelling status
  Task 16  import of both passes (vocabulary checks, hidden-repeat de-aliasing), pass digests
  Task 17  pass agreement, blind-repeat consistency, founder review selection and workbook
  Task 18  founder decision import and the K3 assembly of one round's labels
  Task 19  labels.csv, read_labels, labels_manifest.json and the round report

The blind queue (context pack §6) carries exactly case_id, borrower_name_raw, lender_names_raw, city,
state -- no baseline flag, no stratum, no score, no address, no date. A round's cases are ordered by a
seeded sha256 rank and cut into chunks of `chunk_size` (one labeller agent per chunk per pass, K14).
Inside every chunk round-half-up(repeat_fraction * chunk rows) cases are repeated under
repeat_alias(seed, round, case_id), which cannot be told apart from a real case_id, and the chunk is
re-ordered by rank. The private key maps every queue row back to its case, stratum and design weight.
"""
from __future__ import annotations

import csv
import math
from dataclasses import dataclass
from pathlib import Path

import pandas as pd

from ucc_ml.config import RunConfig, artefact_paths
from ucc_ml.contracts import LABEL_DISCLOSURE, LABELLING_ROUNDS, canonical_json, check_column, is_sha256_hex, sha256_hex
from ucc_ml.provenance import sha256_file
from ucc_ml.sampling import draw_rank

QUEUE_COLUMNS: tuple[str, ...] = ("case_id", "borrower_name_raw", "lender_names_raw", "city", "state")
KEY_COLUMNS: tuple[str, ...] = (
    "queue_case_id", "case_id", "is_repeat", "part", "sampling_stratum", "N_h", "pool_h", "n_h",
    "inclusion_probability", "baseline_qualifies", "queue_rank",
)
ROUND_CASE_COLUMNS: tuple[str, ...] = (
    "case_id", "borrower_name_raw", "lender_names_raw", "borrower_city", "borrower_state",
    "sampling_stratum", "N_h", "pool_h", "n_h", "inclusion_probability", "baseline_qualifies",
)
LENDER_JOIN = " | "
PASSES: tuple[str, ...] = ("a", "b")


@dataclass(frozen=True)
class RoundPaths:
    """Every file of one labelling round. Private files live under ml/data; digests under docs/data/ml."""
    round_name: str
    cases: Path
    manifest: Path
    queue_dir: Path
    key: Path
    report: Path
    labels_dir: Path
    raw_dir: Path
    structured_dir: Path
    briefs_dir: Path
    passes_dir: Path
    agreement: Path
    workbook: Path
    review_manifest: Path
    founder_decisions: Path
    preregistration: Path
    passes_digest: Path
    founder_digest: Path

    def queue_part(self, part: int) -> Path:
        return self.queue_dir / f"queue_{self.round_name}_part_{part:03d}.csv"

    def brief(self, part: int) -> Path:
        return self.briefs_dir / f"brief_{self.round_name}_part_{part:03d}.md"

    def structured_output(self, pass_letter: str, part: int) -> Path:
        return self.structured_dir / f"labeller_output_{self.round_name}_pass_{pass_letter}_part_{part:03d}.json"

    def raw_output(self, pass_letter: str, part: int) -> Path:
        return self.raw_dir / f"labeller_output_{self.round_name}_pass_{pass_letter}_part_{part:03d}.csv"

    def pass_file(self, pass_letter: str) -> Path:
        return self.passes_dir / f"pass_{pass_letter}_{self.round_name}.csv"


def round_paths(cfg: RunConfig, round_name: str) -> RoundPaths:
    if round_name not in LABELLING_ROUNDS:
        raise ValueError(f"unknown labelling round {round_name!r}; expected one of {LABELLING_ROUNDS}")
    paths = artefact_paths(cfg)
    if round_name == "pilot_v1":
        cases, manifest, queue_dir = paths.pilot_cases, paths.pilot_manifest, cfg.path("pilot_dir")
        report = queue_dir / "pilot_report.json"
    else:
        cases, manifest, queue_dir = paths.main_cases, paths.main_manifest, cfg.path("main_round_dir")
        report = queue_dir / "main_report.json"
    labels_dir = cfg.path("labels_dir")
    public = paths.public_data_dir
    return RoundPaths(
        round_name=round_name, cases=cases, manifest=manifest, queue_dir=queue_dir,
        key=queue_dir / f"queue_key_{round_name}.csv", report=report, labels_dir=labels_dir,
        raw_dir=labels_dir / "raw", structured_dir=labels_dir / "raw" / "structured",
        briefs_dir=labels_dir / "briefs", passes_dir=labels_dir / "passes",
        agreement=labels_dir / f"agreement_{round_name}.json",
        workbook=labels_dir / f"founder_review_{round_name}.xlsx",
        review_manifest=labels_dir / f"founder_review_{round_name}.json",
        founder_decisions=labels_dir / "raw" / f"founder_decisions_{round_name}.csv",
        preregistration=public / f"{round_name}.sha256",
        passes_digest=public / f"{round_name}_passes.sha256",
        founder_digest=public / f"{round_name}_founder_review.sha256",
    )


def repeat_alias(seed: int, round_name: str, case_id: str) -> str:
    return sha256_hex(canonical_json(["repeat", round_name, seed, case_id]))


def queue_rank(seed: int, round_name: str, queue_case_id: str) -> str:
    return sha256_hex(canonical_json(["queue_order", round_name, seed, queue_case_id]))


def _round_half_up(value: float) -> int:
    return int(math.floor(value + 0.5))


def _text(value) -> str:
    return value if isinstance(value, str) else ""


def build_blind_queue(round_cases: pd.DataFrame, round_name: str, repeat_fraction: float, seed: int,
                      chunk_size: int) -> tuple[list[pd.DataFrame], pd.DataFrame]:
    """Return (chunks, key): one blind frame per part (QUEUE_COLUMNS) and the private key (KEY_COLUMNS)."""
    if round_name not in LABELLING_ROUNDS:
        raise ValueError(f"unknown labelling round {round_name!r}")
    missing = set(ROUND_CASE_COLUMNS) - set(round_cases.columns)
    if missing:
        raise ValueError(f"round cases lack columns {sorted(missing)}")
    if not round_cases.case_id.is_unique:
        raise ValueError("round cases repeat a case_id")
    if chunk_size < 1:
        raise ValueError("chunk_size must be >= 1")
    design = round_cases.set_index("case_id")
    order = sorted(design.index, key=lambda cid: queue_rank(seed, round_name, cid))
    key_rows: list[dict] = []
    for start in range(0, len(order), chunk_size):
        part = start // chunk_size + 1
        members = order[start:start + chunk_size]
        n_repeats = _round_half_up(repeat_fraction * len(members))
        repeats = sorted(members, key=lambda cid: draw_rank(seed, f"{round_name}_repeat", cid))[:n_repeats]
        entries = [(cid, cid, False) for cid in members]
        entries += [(cid, repeat_alias(seed, round_name, cid), True) for cid in repeats]
        for case_id, queue_case_id, is_repeat in entries:
            row = design.loc[case_id]
            key_rows.append({
                "queue_case_id": queue_case_id, "case_id": case_id, "is_repeat": is_repeat, "part": part,
                "sampling_stratum": row.sampling_stratum, "N_h": int(row.N_h), "pool_h": int(row.pool_h),
                "n_h": int(row.n_h), "inclusion_probability": float(row.inclusion_probability),
                "baseline_qualifies": bool(row.baseline_qualifies),
                "queue_rank": queue_rank(seed, round_name, queue_case_id),
            })
    key = pd.DataFrame(key_rows, columns=list(KEY_COLUMNS))
    key = key.sort_values(["part", "queue_rank"], kind="mergesort").reset_index(drop=True)
    chunks = []
    for _, rows in key.groupby("part", sort=True):
        chunks.append(pd.DataFrame({
            "case_id": rows.queue_case_id.tolist(),
            "borrower_name_raw": [design.at[c, "borrower_name_raw"] for c in rows.case_id],
            "lender_names_raw": [LENDER_JOIN.join(design.at[c, "lender_names_raw"]) for c in rows.case_id],
            "city": [_text(design.at[c, "borrower_city"]) for c in rows.case_id],
            "state": [_text(design.at[c, "borrower_state"]) for c in rows.case_id],
        }, columns=list(QUEUE_COLUMNS)))
    return chunks, key


def write_csv(df: pd.DataFrame, path: Path) -> str:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=False, encoding="utf-8", lineterminator="\n", quoting=csv.QUOTE_MINIMAL)
    return sha256_file(path)


def read_exact_csv(path: Path, columns: tuple[str, ...]) -> pd.DataFrame:
    """Every value as str ('' for empty); refuses a header that is not exactly `columns`."""
    path = Path(path)
    df = pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8")
    if list(df.columns) != list(columns):
        raise ValueError(f"{path}: columns {list(df.columns)} != {list(columns)}")
    return df


_BOOL_TEXT = {"True": True, "False": False}


def parse_bool_column(path: Path, df: pd.DataFrame, column: str) -> None:
    check_column(path, df, column, lambda v: v in _BOOL_TEXT)
    df[column] = pd.Series([_BOOL_TEXT[v] for v in df[column]], index=df.index, dtype=bool)


def read_queue(path: Path) -> pd.DataFrame:
    df = read_exact_csv(path, QUEUE_COLUMNS)
    check_column(path, df, "case_id", is_sha256_hex)
    return df


def read_key(path: Path) -> pd.DataFrame:
    df = read_exact_csv(path, KEY_COLUMNS)
    for column in ("queue_case_id", "case_id"):
        check_column(path, df, column, is_sha256_hex)
    for column in ("part", "N_h", "pool_h", "n_h"):
        check_column(path, df, column, lambda v: v.isdigit())
        df[column] = df[column].astype(int)
    df["inclusion_probability"] = df.inclusion_probability.astype(float)
    for column in ("is_repeat", "baseline_qualifies"):
        parse_bool_column(path, df, column)
    return df


def labels_exist(rp: RoundPaths) -> bool:
    return any(rp.raw_dir.glob(f"labeller_output_{rp.round_name}_pass_*_part_*.csv"))


def write_digest_file(path: Path, files: list[Path], comments: list[str]) -> str:
    """`<sha256>  <file name>` for every file, then `# ` comment lines; returns the digest file's sha256."""
    path = Path(path)
    lines = [f"{sha256_file(f)}  {Path(f).name}" for f in files]
    lines += [f"# {c}" if c else "#" for c in comments]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return sha256_file(path)


def write_preregistration(rp: RoundPaths, part_paths: list[Path]) -> str:
    comments = [
        f"Committed BEFORE any {rp.round_name} label exists (contract K13). The queue chunks are exactly what",
        "the blind labellers see: case_id, borrower name, lender names, city, state -- nothing else. The key",
        "holds the strata, N_h, inclusion probabilities and repeat aliases and stays private. Committing every",
        "hash first proves the sample, the design weights and the repeat structure were fixed in advance.",
        f"Labels for this round will be {LABEL_DISCLOSURE}.",
    ]
    return write_digest_file(rp.preregistration, [rp.cases, rp.manifest, rp.key, *part_paths], comments)


# ============================================================================= Task 15
# Labeller briefs (contract K14), structured-output validation, raw labeller CSVs, labelling status.
import io  # noqa: E402
import re  # noqa: E402
from collections import Counter  # noqa: E402

from ucc_ml.contracts import LABELS, REASON_CODES  # noqa: E402

BRIEF_MARKER = "\n=== BRIEF ===\n"
RAW_OUTPUT_COLUMNS: tuple[str, ...] = ("case_id", "label", "reason_code", "reason")
LABELLER_OUTPUT_SCHEMA: dict = {
    "type": "object",
    "additionalProperties": False,
    "required": ["rows"],
    "properties": {
        "rows": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": list(RAW_OUTPUT_COLUMNS),
                "properties": {
                    "case_id": {"type": "string", "pattern": "^[0-9a-f]{64}$"},
                    "label": {"type": "string", "enum": list(LABELS)},
                    "reason_code": {"type": "string",
                                    "enum": [code for label in LABELS for code in REASON_CODES[label]]},
                    "reason": {"type": "string", "minLength": 1, "maxLength": 300},
                },
            },
        },
    },
}
_URL = re.compile(r"(https?://|www\.)", re.IGNORECASE)


def policy_is_frozen(policy_text: str) -> bool:
    return any(line.strip() == "status: FROZEN" for line in policy_text.splitlines())


def render_labeller_brief(prompt_text: str, policy_text: str, chunk: pd.DataFrame) -> str:
    """The whole prompt of one blind labeller agent: the brief, the policy verbatim, the chunk as CSV.

    Pass A and pass B receive byte-identical briefs for the same chunk; nothing tells an agent which
    pass it is, which round it is in, or anything about the sample design."""
    if prompt_text.count(BRIEF_MARKER) != 1:
        raise ValueError("the labeller prompt must contain exactly one '=== BRIEF ===' line")
    if list(chunk.columns) != list(QUEUE_COLUMNS):
        raise ValueError(f"a queue chunk has exactly the columns {list(QUEUE_COLUMNS)}")
    brief = prompt_text.split(BRIEF_MARKER, 1)[1].strip("\n")
    buffer = io.StringIO()
    chunk.to_csv(buffer, index=False, lineterminator="\n", quoting=csv.QUOTE_MINIMAL)
    return (f"{brief}\n\n---\n\n# The label policy\n\n{policy_text.strip()}\n\n---\n\n"
            f"# Your queue chunk ({len(chunk)} rows)\n\n```csv\n{buffer.getvalue()}```\n")


def rows_from_structured_output(obj) -> pd.DataFrame:
    """The labeller's structured output ({"rows": [...]}) as RAW_OUTPUT_COLUMNS; anything else is refused."""
    if not isinstance(obj, dict) or set(obj) != {"rows"} or not isinstance(obj["rows"], list):
        raise ValueError('structured output must be an object with exactly one key, "rows", holding a list')
    for i, row in enumerate(obj["rows"], start=1):
        if not isinstance(row, dict) or set(row) != set(RAW_OUTPUT_COLUMNS):
            raise ValueError(f"row {i}: keys must be exactly {list(RAW_OUTPUT_COLUMNS)}")
        if not all(isinstance(row[k], str) for k in RAW_OUTPUT_COLUMNS):
            raise ValueError(f"row {i}: every value must be a string")
    return pd.DataFrame([[row[k] for k in RAW_OUTPUT_COLUMNS] for row in obj["rows"]],
                        columns=list(RAW_OUTPUT_COLUMNS), dtype=object)


def validate_labeller_rows(rows: pd.DataFrame, expected_case_ids: list[str], reason_max_chars: int) -> list[str]:
    """Every problem with one pass's rows for one chunk; an empty list means the rows are acceptable."""
    if list(rows.columns) != list(RAW_OUTPUT_COLUMNS):
        return [f"columns {list(rows.columns)} != {list(RAW_OUTPUT_COLUMNS)}"]
    problems: list[str] = []
    got = [v if isinstance(v, str) else "" for v in rows.case_id]
    expected = list(expected_case_ids)
    if got != expected:
        missing = [c for c in expected if c not in set(got)]
        unexpected = [c for c in got if c not in set(expected)]
        duplicated = sorted(c for c, n in Counter(got).items() if n > 1)
        problems.append(f"case_id sequence differs from the queue chunk ({len(got)} rows, {len(expected)} expected; "
                        f"missing {missing[:3]}, unexpected {unexpected[:3]}, duplicated {duplicated[:3]})")
    for i, (label, code, reason) in enumerate(zip(rows.label, rows.reason_code, rows.reason), start=1):
        label = label if isinstance(label, str) else ""
        code = code if isinstance(code, str) else ""
        reason = reason if isinstance(reason, str) else ""
        if label not in LABELS:
            problems.append(f"row {i}: label {label!r} is not one of {list(LABELS)}")
        elif code not in REASON_CODES[label]:
            problems.append(f"row {i}: reason_code {code!r} is not valid for {label}")
        if not reason.strip():
            problems.append(f"row {i}: reason is blank")
        if len(reason) > reason_max_chars:
            problems.append(f"row {i}: reason has {len(reason)} characters (max {reason_max_chars})")
        if "\n" in reason or "\r" in reason:
            problems.append(f"row {i}: reason contains a line break")
        if _URL.search(reason):
            problems.append(f"row {i}: reason contains a URL")
    return problems


def write_raw_output(rows: pd.DataFrame, path: Path, expected_case_ids: list[str], reason_max_chars: int) -> str:
    """Write one pass's answer for one chunk. Refuses invalid rows; never overwrites an existing file."""
    path = Path(path)
    if path.exists():
        raise FileExistsError(f"{path} already exists; raw labeller output is never overwritten")
    problems = validate_labeller_rows(rows, expected_case_ids, reason_max_chars)
    if problems:
        raise ValueError(f"{path.name}: " + "; ".join(problems[:10]))
    return write_csv(rows[list(RAW_OUTPUT_COLUMNS)], path)


def read_raw_output(path: Path) -> pd.DataFrame:
    return read_exact_csv(path, RAW_OUTPUT_COLUMNS)


def queue_parts(key: pd.DataFrame) -> list[int]:
    return sorted({int(p) for p in key.part})


def labelling_status(rp: RoundPaths, parts: list[int]) -> dict[str, dict[str, list[int]]]:
    return {letter: {"present": [p for p in parts if rp.raw_output(letter, p).exists()],
                     "missing": [p for p in parts if not rp.raw_output(letter, p).exists()]}
            for letter in PASSES}


# ============================================================================= Task 16
# Import both blind passes: vocabulary checks, hidden-repeat de-aliasing, pass files, pass digests.
from collections.abc import Mapping  # noqa: E402
from datetime import datetime, timezone  # noqa: E402

from ucc_ml.contracts import PASS_LABELLERS  # noqa: E402

PASS_COLUMNS: tuple[str, ...] = (
    "queue_case_id", "case_id", "is_repeat", "part", "label", "reason_code", "reason", "labeller_id",
    "labelled_at", "sampling_stratum", "inclusion_probability", "labelling_round",
)


def file_mtime_iso(path: Path) -> str:
    return datetime.fromtimestamp(Path(path).stat().st_mtime, tz=timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def import_pass(key: pd.DataFrame, chunks: Mapping[int, pd.DataFrame], raw_outputs: Mapping[int, pd.DataFrame],
                pass_letter: str, round_name: str, labelled_at: Mapping[int, str],
                reason_max_chars: int) -> pd.DataFrame:
    """One blind pass over a whole round as PASS_COLUMNS rows, every hidden repeat resolved to its case."""
    if pass_letter not in PASSES:
        raise ValueError(f"unknown pass {pass_letter!r}")
    if round_name not in LABELLING_ROUNDS:
        raise ValueError(f"unknown labelling round {round_name!r}")
    parts = queue_parts(key)
    missing = [p for p in parts if p not in raw_outputs]
    if missing:
        raise ValueError(f"{round_name} pass {pass_letter}: labeller output missing for part(s) {missing}")
    extra = sorted(set(raw_outputs) - set(parts))
    if extra:
        raise ValueError(f"{round_name} pass {pass_letter}: labeller output for unknown part(s) {extra}")
    frames = []
    for part in parts:
        queue_ids = key.loc[key.part == part, "queue_case_id"].tolist()
        if part not in chunks or chunks[part].case_id.tolist() != queue_ids:
            raise ValueError(f"{round_name} part {part:03d}: the queue chunk and the private key disagree")
        problems = validate_labeller_rows(raw_outputs[part], queue_ids, reason_max_chars)
        if problems:
            raise ValueError(f"{round_name} pass {pass_letter} part {part:03d}: " + "; ".join(problems[:10]))
        rows = raw_outputs[part][list(RAW_OUTPUT_COLUMNS)].rename(columns={"case_id": "queue_case_id"}).copy()
        rows["part"] = part
        rows["labelled_at"] = labelled_at[part]
        frames.append(rows)
    answers = pd.concat(frames, ignore_index=True)
    design = key[["queue_case_id", "case_id", "is_repeat", "sampling_stratum", "inclusion_probability"]]
    merged = answers.merge(design, on="queue_case_id", how="left", validate="one_to_one")
    merged["labeller_id"] = PASS_LABELLERS[pass_letter]
    merged["labelling_round"] = round_name
    return merged[list(PASS_COLUMNS)]


def read_pass_file(path: Path) -> pd.DataFrame:
    path = Path(path)
    df = read_exact_csv(path, PASS_COLUMNS)
    for column in ("queue_case_id", "case_id"):
        check_column(path, df, column, is_sha256_hex)
    check_column(path, df, "labelling_round", lambda v: v in LABELLING_ROUNDS)
    check_column(path, df, "labeller_id", lambda v: v in PASS_LABELLERS.values())
    check_column(path, df, "label", lambda v: v in LABELS)
    check_column(path, df, "part", lambda v: v.isdigit())
    df["part"] = df.part.astype(int)
    df["inclusion_probability"] = df.inclusion_probability.astype(float)
    parse_bool_column(path, df, "is_repeat")
    return df


def import_round_passes(rp: RoundPaths, reason_max_chars: int) -> dict[str, pd.DataFrame]:
    """Import both passes of a round from its key, its chunks and every raw labeller output."""
    key = read_key(rp.key)
    parts = queue_parts(key)
    status = labelling_status(rp, parts)
    missing = {letter: s["missing"] for letter, s in status.items() if s["missing"]}
    if missing:
        raise ValueError(f"{rp.round_name}: labeller output missing for {missing}")
    chunks = {p: read_queue(rp.queue_part(p)) for p in parts}
    passes = {}
    for letter in PASSES:
        raws = {p: read_raw_output(rp.raw_output(letter, p)) for p in parts}
        stamps = {p: file_mtime_iso(rp.raw_output(letter, p)) for p in parts}
        passes[letter] = import_pass(key, chunks, raws, letter, rp.round_name, stamps, reason_max_chars)
    return passes


def write_passes_digest(rp: RoundPaths, parts: list[int]) -> str:
    files = [rp.raw_output(letter, p) for letter in PASSES for p in parts]
    files += [rp.pass_file(letter) for letter in PASSES]
    comments = [
        f"The {rp.round_name} blind labeller outputs, frozen before the founder review (contract K14): two",
        "independent tool-less Claude passes, one agent per queue chunk per pass; the orchestrator wrote each",
        "raw CSV from the agent's structured output. The last two lines are the imported pass files.",
        f"Labels built from these passes are {LABEL_DISCLOSURE}.",
    ]
    return write_digest_file(rp.passes_digest, files, comments)


# ============================================================================= Task 17
# Pass agreement, blind-repeat consistency, the founder review selection and the founder workbook.
from ucc_ml.contracts import STRATA  # noqa: E402
from ucc_ml.provenance import finite_or_none  # noqa: E402

REVIEW_COLUMNS: tuple[str, ...] = ("case_id", "review_type", "split", "sampling_stratum", "review_rank")
WORKBOOK_SHEET = "REVIEW THESE"
GUIDE_SHEET = "HOW TO REVIEW"
WORKBOOK_COLUMNS: tuple[str, ...] = (
    "row", "case_id", "borrower_name_raw", "lender_names_raw", "city", "state",
    "pass_a_label", "pass_a_reason_code", "pass_a_reason",
    "pass_b_label", "pass_b_reason_code", "pass_b_reason",
    "founder_label", "founder_note",
)
FOUNDER_GUIDE: tuple[str, ...] = (
    "Founder review of blind Claude labels -- label_policy_v1 (ml/specs/label_policy_v1.md)",
    "",
    "Each row is one case: the borrower name, its lender names, city and state, and what two independent",
    "blind Claude passes decided. Rows where the passes disagree are mixed with a random audit sample of rows",
    "where they agree. Nothing else is shown on purpose: no rule result, no sample design, no model output.",
    "",
    "For every row choose founder_label: RELEVANT, NOT_RELEVANT or INSUFFICIENT_EVIDENCE.",
    "  - The passes disagree: your label decides the case. founder_note is required (one sentence: why).",
    "  - The passes agree and so do you: choose the same label. founder_note is optional.",
    "  - The passes agree and you do not: choose your label. founder_note is required (why they were wrong).",
    "A blank founder_label means not reviewed. Labels are not written while any disagreement is blank.",
    "",
    "Do not look anything up: no Google, no Secretary of State. Decide from the names and city/state shown.",
    "Only fill founder_label and founder_note. Do not add, delete or re-type rows or other columns.",
)


def _originals(pass_frame: pd.DataFrame) -> pd.DataFrame:
    originals = pass_frame[~pass_frame.is_repeat].set_index("case_id")
    if not originals.index.is_unique:
        raise ValueError("a pass lists the same original case more than once")
    return originals


def _agreement(n: int, agreed: int) -> dict:
    return {"n": int(n), "agreed": int(agreed), "rate": finite_or_none(agreed / n) if n else None}


def repeat_consistency(pass_frame: pd.DataFrame) -> dict:
    """Share of one pass's hidden repeats that got the same label as that pass gave the original case."""
    originals = _originals(pass_frame).label
    repeats = pass_frame[pass_frame.is_repeat]
    consistent = sum(1 for case_id, label in zip(repeats.case_id, repeats.label) if originals.get(case_id) == label)
    n = len(repeats)
    return {"n": int(n), "consistent": int(consistent), "rate": finite_or_none(consistent / n) if n else None}


def _insufficient_share(frame: pd.DataFrame) -> dict:
    shares = {}
    for stratum in STRATA:
        mask = frame.sampling_stratum == stratum
        if mask.any():
            shares[stratum] = finite_or_none(float((frame.label[mask] == "INSUFFICIENT_EVIDENCE").mean()))
    return shares


def agreement_report(pass_a: pd.DataFrame, pass_b: pd.DataFrame, round_name: str) -> dict:
    """Pass agreement overall and by stratum, label counts, confusion, repeat consistency, disagreements."""
    a, b = _originals(pass_a), _originals(pass_b)
    if set(a.index) != set(b.index):
        raise ValueError(f"{round_name}: pass a and pass b cover different cases")
    b = b.loc[a.index]
    agreed = a.label == b.label
    by_stratum = {s: _agreement(int((a.sampling_stratum == s).sum()), int((agreed & (a.sampling_stratum == s)).sum()))
                  for s in STRATA if (a.sampling_stratum == s).any()}
    return {
        "round": round_name,
        "disclosure": LABEL_DISCLOSURE,
        "pass_agreement": {**_agreement(len(a), int(agreed.sum())), "by_stratum": by_stratum},
        "label_counts": {f"pass_{x}": {label: int((frame.label == label).sum()) for label in LABELS}
                         for x, frame in (("a", a), ("b", b))},
        "insufficient_share_by_stratum": {"pass_a": _insufficient_share(a), "pass_b": _insufficient_share(b)},
        "confusion": dict(sorted(Counter(f"{x}->{y}" for x, y in zip(a.label, b.label)).items())),
        "repeat_consistency": {"pass_a": repeat_consistency(pass_a), "pass_b": repeat_consistency(pass_b)},
        "disagreements": sorted(a.index[~agreed]),
    }


def select_founder_review(pass_a: pd.DataFrame, pass_b: pd.DataFrame, splits: pd.DataFrame, per_split_stratum: int,
                          seed: int, round_name: str) -> pd.DataFrame:
    """Every disagreement plus, in each (split, stratum), the `per_split_stratum` agreed cases with the
    smallest seeded rank; ordered by a seeded rank so audits and disagreements are interleaved."""
    a, b = _originals(pass_a), _originals(pass_b)
    if set(a.index) != set(b.index):
        raise ValueError(f"{round_name}: pass a and pass b cover different cases")
    split_of = splits.set_index("case_id").split
    unknown = sorted(set(a.index) - set(split_of.index))
    if unknown:
        raise ValueError(f"{round_name}: {len(unknown)} labelled case(s) have no split, e.g. {unknown[:3]}")
    frame = pd.DataFrame({
        "case_id": list(a.index),
        "split": [split_of[c] for c in a.index],
        "sampling_stratum": a.sampling_stratum.tolist(),
        "agreed": (a.label == b.loc[a.index].label).tolist(),
    })
    chosen = [(c, "disagreement") for c in frame.case_id[~frame.agreed]]
    for _, group in frame[frame.agreed].groupby(["split", "sampling_stratum"], sort=True):
        ranked = sorted(group.case_id, key=lambda c: draw_rank(seed, f"{round_name}_founder_audit", c))
        chosen += [(c, "audit") for c in ranked[:per_split_stratum]]
    rows = frame.set_index("case_id")
    review = pd.DataFrame({
        "case_id": [c for c, _ in chosen],
        "review_type": [t for _, t in chosen],
        "split": [rows.at[c, "split"] for c, _ in chosen],
        "sampling_stratum": [rows.at[c, "sampling_stratum"] for c, _ in chosen],
        "review_rank": [sha256_hex(canonical_json(["founder_review", round_name, seed, c])) for c, _ in chosen],
    }, columns=list(REVIEW_COLUMNS))
    return review.sort_values("review_rank", kind="mergesort").reset_index(drop=True)


def build_founder_workbook(review: pd.DataFrame, pass_a: pd.DataFrame, pass_b: pd.DataFrame,
                           round_cases: pd.DataFrame, path: Path) -> str:
    """Write the founder's review workbook and return its sha256 as issued.

    Columns are WORKBOOK_COLUMNS: names, city/state, both passes' label, reason code and reason, and the
    empty founder_label / founder_note columns. No baseline flag, stratum, split, score or review type."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    a, b = _originals(pass_a), _originals(pass_b)
    cases = round_cases.set_index("case_id")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = WORKBOOK_SHEET
    sheet.append(list(WORKBOOK_COLUMNS))
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for number, case_id in enumerate(review.case_id, start=1):
        case, answer_a, answer_b = cases.loc[case_id], a.loc[case_id], b.loc[case_id]
        sheet.append([number, case_id, case.borrower_name_raw, LENDER_JOIN.join(case.lender_names_raw),
                      _text(case.borrower_city), _text(case.borrower_state),
                      answer_a.label, answer_a.reason_code, answer_a.reason,
                      answer_b.label, answer_b.reason_code, answer_b.reason, None, None])
    last_row = len(review) + 1
    label_letter = sheet.cell(row=1, column=WORKBOOK_COLUMNS.index("founder_label") + 1).column_letter
    note_letter = sheet.cell(row=1, column=WORKBOOK_COLUMNS.index("founder_note") + 1).column_letter
    validation = DataValidation(type="list", formula1='"' + ",".join(LABELS) + '"', allow_blank=True)
    validation.error = "founder_label must be RELEVANT, NOT_RELEVANT or INSUFFICIENT_EVIDENCE"
    sheet.add_data_validation(validation)
    if last_row >= 2:
        validation.add(f"{label_letter}2:{label_letter}{last_row}")
        sheet.conditional_formatting.add(
            f"A2:{note_letter}{last_row}",
            FormulaRule(formula=[f'${label_letter}2=""'], fill=PatternFill("solid", fgColor="FFF3CD")))
    for column, width in enumerate([6, 18, 34, 40, 14, 6, 22, 26, 40, 22, 26, 40, 24, 40], start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=column).column_letter].width = width
    sheet.freeze_panes = "C2"
    guide = workbook.create_sheet(GUIDE_SHEET, 0)
    for line in FOUNDER_GUIDE:
        guide.append([line])
    guide.column_dimensions["A"].width = 110
    workbook.active = 1
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return sha256_file(path)
