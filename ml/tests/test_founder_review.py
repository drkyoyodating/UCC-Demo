"""The founder's decisions become founder_confirmed / founder_adjudicated rows (K3); undecided disagreements refuse."""
import json

import pandas as pd
import pytest

from conftest import fill_founder_workbook, founder_answers, imported_pilot_repo, reviewed_pilot_repo


@pytest.fixture(scope="module")
def issued_review(tmp_path_factory):
    from ucc_ml.cli import main
    from ucc_ml.config import load_config
    from ucc_ml.labeling import read_pass_file, round_paths

    cfg, disagreements = imported_pilot_repo(tmp_path_factory.mktemp("repo"))
    assert main(["review-workbook", "--config", str(cfg), "--round", "pilot_v1"]) == 0
    rp = round_paths(load_config(cfg), "pilot_v1")
    manifest = json.loads(rp.review_manifest.read_text())
    return cfg, rp, manifest, read_pass_file(rp.pass_file("a")), read_pass_file(rp.pass_file("b"))


def _filled_rows(issued_review, tmp_path, answers):
    import shutil

    from ucc_ml.labeling import read_founder_workbook

    _, rp, *_ = issued_review
    tmp_path.mkdir(parents=True, exist_ok=True)
    copy = tmp_path / rp.workbook.name
    shutil.copy(rp.workbook, copy)
    fill_founder_workbook(copy, answers)
    return read_founder_workbook(copy)


def test_import_founder_review_classifies_every_answer(issued_review, tmp_path):
    from ucc_ml.labeling import FOUNDER_DECISION_COLUMNS, import_founder_review

    cfg, rp, manifest, pass_a, pass_b = issued_review
    answers = founder_answers(cfg, "pilot_v1")
    rows = _filled_rows(issued_review, tmp_path, answers)
    decisions, summary = import_founder_review(rows, manifest, pass_a, pass_b, "2026-09-13T09:00:00Z")
    assert list(decisions.columns) == list(FOUNDER_DECISION_COLUMNS)
    kinds = decisions.set_index("case_id").decision
    assert all(kinds[c] == "adjudicated" for c in manifest["disagreements"])
    assert kinds[manifest["audit"][0]] == "confirmed" and kinds[manifest["audit"][1]] == "overturned"
    assert len(decisions) == len(manifest["disagreements"]) + 2           # blank audit rows are not decisions
    assert summary == {
        "disagreements": {"n": 3, "decided": 3, "undecided": 0},
        "audit": {"n_selected": len(manifest["audit"]), "n_audited": 2, "n_confirmed": 1, "n_overturned": 1,
                  "agreement_rate": 0.5},
    }
    partial = _filled_rows(issued_review, tmp_path / "partial", {c: answers[c] for c in manifest["disagreements"][1:]})
    _, partial_summary = import_founder_review(partial, manifest, pass_a, pass_b, "2026-09-13T09:00:00Z")
    assert partial_summary["disagreements"] == {"n": 3, "decided": 2, "undecided": 1}


@pytest.mark.parametrize("change, message", [
    ("unknown_label", "is not one of"),
    ("disagreement_without_note", "a disagreement decision needs a founder_note"),
    ("overturn_without_note", "needs a founder_note"),
    ("extra_row", "differ from the issued review"),
    ("missing_row", "differ from the issued review"),
])
def test_import_founder_review_rejects_bad_answers(issued_review, tmp_path, change, message):
    from ucc_ml.labeling import import_founder_review

    cfg, rp, manifest, pass_a, pass_b = issued_review
    answers = founder_answers(cfg, "pilot_v1")
    first_disagreement, second_audit = manifest["disagreements"][0], manifest["audit"][1]
    if change == "unknown_label":
        answers[first_disagreement] = ("MAYBE", "unsure")
    elif change == "disagreement_without_note":
        answers[first_disagreement] = ("RELEVANT", "")
    elif change == "overturn_without_note":
        answers[second_audit] = (answers[second_audit][0], "")
    rows = _filled_rows(issued_review, tmp_path, answers)
    if change == "extra_row":
        rows = pd.concat([rows, rows.tail(1).assign(case_id="f" * 64)], ignore_index=True)
    elif change == "missing_row":
        rows = rows.iloc[1:].reset_index(drop=True)
    with pytest.raises(ValueError, match=message):
        import_founder_review(rows, manifest, pass_a, pass_b, "2026-09-13T09:00:00Z")


def test_read_founder_workbook_refuses_a_changed_header(issued_review, tmp_path):
    import shutil

    from openpyxl import load_workbook

    from ucc_ml.labeling import WORKBOOK_SHEET, read_founder_workbook

    _, rp, *_ = issued_review
    copy = tmp_path / "changed.xlsx"
    shutil.copy(rp.workbook, copy)
    workbook = load_workbook(copy)
    workbook[WORKBOOK_SHEET].cell(row=1, column=13, value="label")
    workbook.save(copy)
    with pytest.raises(ValueError, match="header"):
        read_founder_workbook(copy)


def _pass(letter, rows):
    from ucc_ml.contracts import PASS_LABELLERS
    from ucc_ml.labeling import PASS_COLUMNS

    records = [{
        "queue_case_id": "e" * 64 if repeat else case_id, "case_id": case_id, "is_repeat": repeat, "part": 1,
        "label": label, "reason_code": code, "reason": f"pass {letter}: {code}", "labeller_id": PASS_LABELLERS[letter],
        "labelled_at": f"2026-09-12T10:00:0{1 if letter == 'a' else 2}Z", "sampling_stratum": "CO:accepted",
        "inclusion_probability": 0.5, "labelling_round": "pilot_v1",
    } for case_id, label, code, repeat in rows]
    return pd.DataFrame(records, columns=list(PASS_COLUMNS))


K = ["1" * 64, "2" * 64, "3" * 64, "4" * 64]
PASS_A = _pass("a", [(K[0], "RELEVANT", "R_BORROWER_EQUIPMENT_WORD", False), (K[1], "RELEVANT", "R_BORROWER_TRADE_WORD", False),
                     (K[2], "RELEVANT", "R_LENDER_NAMED_MAKER", False), (K[3], "RELEVANT", "R_BORROWER_TRADE_WORD", False),
                     (K[0], "RELEVANT", "R_BORROWER_EQUIPMENT_WORD", True)])
PASS_B = _pass("b", [(K[0], "RELEVANT", "R_BORROWER_EQUIPMENT_WORD", False), (K[1], "RELEVANT", "R_BORROWER_TRADE_WORD", False),
                     (K[2], "RELEVANT", "R_LENDER_NAMED_MAKER", False), (K[3], "INSUFFICIENT_EVIDENCE", "I_AMBIGUOUS_WORD", False),
                     (K[0], "INSUFFICIENT_EVIDENCE", "I_GENERIC_NAME", True)])


def _decisions(rows):
    from ucc_ml.labeling import FOUNDER_DECISION_COLUMNS

    return pd.DataFrame([dict(zip(FOUNDER_DECISION_COLUMNS, r)) for r in rows], columns=list(FOUNDER_DECISION_COLUMNS))


DECISIONS = _decisions([
    (K[1], "audit", "RELEVANT", "", "confirmed", "2026-09-13T09:00:00Z"),
    (K[2], "audit", "NOT_RELEVANT", "a law office, not a Terex customer", "overturned", "2026-09-13T09:00:00Z"),
    (K[3], "disagreement", "RELEVANT", "CONCRETE is a trade word", "adjudicated", "2026-09-13T09:00:00Z"),
])


def test_assemble_round_labels_produces_the_four_k3_statuses():
    from ucc_ml.contracts import LABEL_COLUMNS, Label
    from ucc_ml.labeling import assemble_round_labels

    labels = assemble_round_labels(PASS_A, PASS_B, DECISIONS, "pilot_v1", "label_policy_v1")
    assert list(labels.columns) == list(LABEL_COLUMNS) and len(labels) == 6
    originals = labels[~labels.is_repeat].set_index("case_id")
    assert originals.adjudication_status.to_dict() == {K[0]: "model_agreed", K[1]: "founder_confirmed",
                                                       K[2]: "founder_adjudicated", K[3]: "founder_adjudicated"}
    assert originals.at[K[0], "labeller_id"] == "claude_blind_pass_a+claude_blind_pass_b"
    assert originals.at[K[0], "labelled_at"] == "2026-09-12T10:00:02Z" and originals.at[K[0], "reason"] == "pass a: R_BORROWER_EQUIPMENT_WORD"
    assert originals.at[K[1], "labeller_id"] == "founder" and originals.at[K[1], "reason_code"] == "R_BORROWER_TRADE_WORD"
    assert originals.loc[K[2], ["label", "reason_code", "reason"]].tolist() == ["NOT_RELEVANT", "ADJUDICATED", "a law office, not a Terex customer"]
    assert originals.at[K[3], "label"] == "RELEVANT" and originals.at[K[3], "reason_code"] == "ADJUDICATED"
    repeats = labels[labels.is_repeat]
    assert repeats.labeller_id.tolist() == ["claude_blind_pass_a", "claude_blind_pass_b"]
    assert set(repeats.adjudication_status) == {"blind_repeat"} and set(repeats.case_id) == {K[0]}
    assert set(labels.labelling_round) == {"pilot_v1"} and set(labels.policy_version) == {"label_policy_v1"}
    for record in labels.to_dict("records"):
        Label.model_validate(record)


def test_assemble_round_labels_refuses_an_undecided_disagreement():
    from ucc_ml.labeling import UndecidedDisagreements, assemble_round_labels

    with pytest.raises(UndecidedDisagreements, match="1 disagreement") as exc:
        assemble_round_labels(PASS_A, PASS_B, DECISIONS.iloc[:2], "pilot_v1", "label_policy_v1")
    assert exc.value.case_ids == [K[3]] and exc.value.round_name == "pilot_v1"
    wrong = _decisions([(K[0], "disagreement", "RELEVANT", "x", "adjudicated", "2026-09-13T09:00:00Z")])
    with pytest.raises(ValueError, match="but the passes agree"):
        assemble_round_labels(PASS_A, PASS_B, pd.concat([DECISIONS, wrong], ignore_index=True), "pilot_v1", "label_policy_v1")


def test_cli_import_founder_review_complete_and_incomplete(tmp_path, capsys):
    import hashlib

    from ucc_ml.config import load_config
    from ucc_ml.labeling import read_founder_decisions, round_paths

    cfg = reviewed_pilot_repo(tmp_path / "complete")
    rp = round_paths(load_config(cfg), "pilot_v1")
    decisions = read_founder_decisions(rp.founder_decisions)
    assert set(decisions.decision) == {"adjudicated", "confirmed", "overturned"}
    lines = rp.founder_digest.read_text().splitlines()
    assert [line.split("  ", 1)[1] for line in lines[:3]] == [
        "founder_review_pilot_v1.xlsx", "founder_review_pilot_v1.json", "founder_decisions_pilot_v1.csv"]
    assert lines[2].split("  ", 1)[0] == hashlib.sha256(rp.founder_decisions.read_bytes()).hexdigest()
    assert "disagreements: 3/3 decided" in capsys.readouterr().out
    reviewed_pilot_repo(tmp_path / "incomplete", decide_all=False)
    assert "INCOMPLETE: 1 disagreement(s)" in capsys.readouterr().out
