"""Agreement, blind-repeat consistency and the founder review workbook: every disagreement plus an audit sample."""
import json

import pandas as pd
import pytest

from conftest import imported_pilot_repo

C = [format(i, "x") * 64 for i in range(1, 7)]


def _frame(rows):
    return pd.DataFrame(rows, columns=["case_id", "is_repeat", "label", "reason_code", "reason", "sampling_stratum"])


PASS_A = _frame([
    (C[0], False, "RELEVANT", "R_BORROWER_EQUIPMENT_WORD", "EXCAVATION", "CO:accepted"),
    (C[1], False, "NOT_RELEVANT", "N_OTHER_INDUSTRY_EXPLICIT", "LAW OFFICE", "CO:rejected"),
    (C[2], False, "INSUFFICIENT_EVIDENCE", "I_GENERIC_NAME", "could be anything", "CT:rejected"),
    (C[3], False, "RELEVANT", "R_LENDER_NAMED_MAKER", "TEREX lender", "CT:accepted"),
    (C[0], True, "RELEVANT", "R_BORROWER_EQUIPMENT_WORD", "EXCAVATION", "CO:accepted"),
    (C[2], True, "NOT_RELEVANT", "N_OTHER_INDUSTRY_EXPLICIT", "a shop", "CT:rejected"),
])
PASS_B = _frame([
    (C[0], False, "RELEVANT", "R_BORROWER_EQUIPMENT_WORD", "EXCAVATION", "CO:accepted"),
    (C[1], False, "INSUFFICIENT_EVIDENCE", "I_AMBIGUOUS_WORD", "OFFICE alone", "CO:rejected"),
    (C[2], False, "INSUFFICIENT_EVIDENCE", "I_GENERIC_NAME", "could be anything", "CT:rejected"),
    (C[3], False, "RELEVANT", "R_LENDER_NAMED_MAKER", "TEREX lender", "CT:accepted"),
    (C[0], True, "RELEVANT", "R_BORROWER_EQUIPMENT_WORD", "EXCAVATION", "CO:accepted"),
    (C[2], True, "INSUFFICIENT_EVIDENCE", "I_GENERIC_NAME", "could be anything", "CT:rejected"),
])


def test_agreement_report_on_a_known_pair():
    from ucc_ml.labeling import agreement_report

    report = agreement_report(PASS_A, PASS_B, "pilot_v1")
    assert report["disclosure"] == "model-labelled, founder-adjudicated"
    assert {k: report["pass_agreement"][k] for k in ("n", "agreed", "rate")} == {"n": 4, "agreed": 3, "rate": 0.75}
    assert report["pass_agreement"]["by_stratum"]["CO:rejected"] == {"n": 1, "agreed": 0, "rate": 0.0}
    assert report["disagreements"] == [C[1]]
    assert report["confusion"] == {"INSUFFICIENT_EVIDENCE->INSUFFICIENT_EVIDENCE": 1,
                                   "NOT_RELEVANT->INSUFFICIENT_EVIDENCE": 1, "RELEVANT->RELEVANT": 2}
    assert report["repeat_consistency"] == {"pass_a": {"n": 2, "consistent": 1, "rate": 0.5},
                                            "pass_b": {"n": 2, "consistent": 2, "rate": 1.0}}
    assert report["label_counts"]["pass_b"] == {"RELEVANT": 2, "NOT_RELEVANT": 0, "INSUFFICIENT_EVIDENCE": 2}
    assert report["insufficient_share_by_stratum"]["pass_b"]["CO:rejected"] == 1.0
    json.dumps(report, allow_nan=False)
    with pytest.raises(ValueError, match="different cases"):
        agreement_report(PASS_A, PASS_B[PASS_B.case_id != C[3]], "pilot_v1")


def _many(n=24, disagree=(0, 5, 7)):
    from ucc_ml.contracts import STRATA, sha256_hex

    ids = [sha256_hex(str(i)) for i in range(n)]
    strata = [STRATA[i % 4] for i in range(n)]
    splits = pd.DataFrame({"case_id": ids, "group_id": ids, "split": [("train", "validation", "test")[i % 3] for i in range(n)]})
    a = _frame([(ids[i], False, "RELEVANT", "R_BORROWER_TRADE_WORD", "CONCRETE", strata[i]) for i in range(n)])
    b = a.copy()
    b.loc[list(disagree), ["label", "reason_code", "reason"]] = ["INSUFFICIENT_EVIDENCE", "I_GENERIC_NAME", "unclear"]
    return a, b, splits, [ids[i] for i in disagree]


def test_select_founder_review_takes_every_disagreement_and_a_per_cell_audit():
    from ucc_ml.labeling import REVIEW_COLUMNS, select_founder_review

    a, b, splits, disagreements = _many()
    review = select_founder_review(a, b, splits, per_split_stratum=2, seed=20260912, round_name="pilot_v1")
    assert list(review.columns) == list(REVIEW_COLUMNS)
    assert sorted(review.case_id[review.review_type == "disagreement"]) == sorted(disagreements)
    audit = review[review.review_type == "audit"]
    assert review.case_id.is_unique and not set(audit.case_id) & set(disagreements)
    per_cell = audit.groupby(["split", "sampling_stratum"]).size()
    assert len(per_cell) == 12 and set(per_cell) <= {1, 2} and int(per_cell.sum()) == 21
    assert review.review_rank.tolist() == sorted(review.review_rank)
    assert review.case_id.tolist() == select_founder_review(a, b, splits, 2, 20260912, "pilot_v1").case_id.tolist()
    assert set(select_founder_review(a, b, splits, 0, 20260912, "pilot_v1").review_type) == {"disagreement"}
    with pytest.raises(ValueError, match="have no split"):
        select_founder_review(a, b, splits.iloc[1:], 2, 20260912, "pilot_v1")


def test_cli_review_workbook_is_blind_to_the_design_and_never_regenerated(tmp_path, capsys):
    from openpyxl import load_workbook

    from ucc_ml.cli import main
    from ucc_ml.config import load_config
    from ucc_ml.dataset import read_frame
    from ucc_ml.labeling import GUIDE_SHEET, WORKBOOK_COLUMNS, WORKBOOK_SHEET, round_paths

    cfg, disagreements = imported_pilot_repo(tmp_path)
    rp = round_paths(load_config(cfg), "pilot_v1")
    capsys.readouterr()
    assert main(["review-workbook", "--config", str(cfg), "--round", "pilot_v1"]) == 0
    out = capsys.readouterr().out
    assert "pass agreement: 37/40 rate=0.925" in out and "disagreements=3" in out
    manifest = json.loads(rp.review_manifest.read_text())
    assert manifest["disagreements"] == sorted(disagreements) and manifest["rows"] == 3 + len(manifest["audit"])
    assert set(manifest["audit_by_split_stratum"]) == {"train"}              # every pilot case is in train (K12)
    assert all(n <= 2 for cell in manifest["audit_by_split_stratum"].values() for n in cell.values())
    workbook = load_workbook(rp.workbook)
    assert workbook.sheetnames == [GUIDE_SHEET, WORKBOOK_SHEET]
    sheet = workbook[WORKBOOK_SHEET]
    rows = list(sheet.iter_rows(values_only=True))
    assert tuple(rows[0]) == WORKBOOK_COLUMNS and len(rows) == manifest["rows"] + 1
    for banned in ("baseline", "stratum", "score", "split", "route", "inclusion", "review_type", "rank"):
        assert not any(banned in str(h) for h in rows[0]), banned
    by_case = {r[1]: r for r in rows[1:]}
    assert set(disagreements) <= set(by_case)
    cases = read_frame(rp.cases).set_index("case_id")
    for case_id, row in by_case.items():
        assert row[2] == cases.at[case_id, "borrower_name_raw"] and row[12] is None and row[13] is None
    assert {row[6] != row[9] for row in rows[1:]} == {True, False}
    assert any(",".join(["RELEVANT", "NOT_RELEVANT", "INSUFFICIENT_EVIDENCE"]) in v.formula1
               for v in sheet.data_validations.dataValidation)
    assert any("Do not look anything up" in (r[0] or "") for r in workbook[GUIDE_SHEET].iter_rows(values_only=True))
    assert json.loads(rp.agreement.read_text())["disagreements"] == sorted(disagreements)
    assert main(["review-workbook", "--config", str(cfg), "--round", "pilot_v1"]) == 1
    assert "REFUSED" in capsys.readouterr().out
