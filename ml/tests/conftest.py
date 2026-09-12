"""Synthetic snapshot builders. No test in ml/tests opens ucc.duckdb or anything under ml/data/.

The column lists are the live schema of ucc.duckdb (context pack §3, verified 2026-09-12); a
snapshot written here has exactly the columns export_snapshot.py would write, so dataset.py is
tested against the real shape with tiny, hand-built rows.
"""
from __future__ import annotations

from pathlib import Path

import duckdb
import pandas as pd
import pytest

from ucc_ml.provenance import sha256_file, write_json

SNAPSHOT_COLUMNS: dict[str, list[str]] = {
    "co_filings": ["transactionid", "masterdocumentid", "transactiontype", "filingtype", "documenttype",
                   "filingdate", "continuation", "terminationflag", "fileid"],
    "co_debtors": ["debtorid", "organizationname", "address1", "city", "state", "zipcode", "fileid",
                   "actiontype", "recordstatus", "efsuniqueid"],
    "co_secured_parties": ["spid", "organizationname", "address1", "city", "state", "zipcode", "fileid",
                           "actiontype", "recordstatus", "assignor"],
    "ct_filings": ["id_lien_flng_nbr", "lien_status", "cd_flng_type", "debtor_nm_bus", "debtor_ad_str1",
                   "debtor_ad_city", "debtor_ad_state", "debtor_ad_zip", "sec_party_nm_bus",
                   "sec_party_ad_str1", "sec_party_ad_city", "sec_party_ad_state", "sec_party_ad_zip",
                   "dt_lapse", "dt_accept"],
    "scope_all": ["region", "fileid", "loan_year", "borrower", "borrower_address", "borrower_city",
                  "borrower_state", "borrower_zip", "lender", "route_a", "route_b"],
}
SOURCE_TABLE_NAMES: dict[str, str] = {
    "co_filings": "filings", "co_debtors": "debtors", "co_secured_parties": "secured_parties",
    "ct_filings": "ct_filings", "scope_all": "scope_all",
}
_BOOLEAN = {("scope_all", "route_a"), ("scope_all", "route_b")}


def _frame(name: str, rows: list[dict]) -> pd.DataFrame:
    cols = SNAPSHOT_COLUMNS[name]
    unknown = {k for r in rows for k in r} - set(cols)
    if unknown:
        raise KeyError(f"{name}: unknown columns {sorted(unknown)}")
    return pd.DataFrame({c: [r.get(c) for r in rows] for c in cols}, columns=cols, dtype=object)


def _typed_select(name: str) -> str:
    parts = []
    for c in SNAPSHOT_COLUMNS[name]:
        typ = "BOOLEAN" if (name, c) in _BOOLEAN else "VARCHAR"
        parts.append(f'"{c}"::{typ} AS "{c}"')
    return ", ".join(parts)


def write_snapshot(snapshot_dir: Path, tables: dict[str, list[dict]]) -> Path:
    """Write every snapshot table (missing ones empty) as Parquet plus a manifest.json."""
    snapshot_dir = Path(snapshot_dir)
    snapshot_dir.mkdir(parents=True, exist_ok=True)
    con = duckdb.connect()
    manifest_tables = {}
    try:
        for name in SNAPSHOT_COLUMNS:
            df = _frame(name, tables.get(name, []))
            con.register("_t", df)
            target = snapshot_dir / f"{name}.parquet"
            con.execute(f"COPY (SELECT {_typed_select(name)} FROM _t) TO '{target.as_posix()}' (FORMAT PARQUET)")
            con.unregister("_t")
            manifest_tables[name] = {"source_table": SOURCE_TABLE_NAMES[name], "columns": SNAPSHOT_COLUMNS[name],
                                     "row_count": len(df), "sha256": sha256_file(target),
                                     "bytes": target.stat().st_size}
    finally:
        con.close()
    write_json(snapshot_dir / "manifest.json", {
        "schema_version": 1, "source_path": "synthetic", "source_bytes": 0, "source_sha256": "synthetic",
        "duckdb_version": duckdb.__version__, "python_version": "test", "exported_at": "2026-01-01T00:00:00Z",
        "tables": manifest_tables,
    })
    return snapshot_dir


def build_source_duckdb(path: Path, tables: dict[str, list[dict]]) -> Path:
    """A tiny ucc.duckdb look-alike with the SOURCE table names, for testing export_snapshot.py."""
    path = Path(path)
    con = duckdb.connect(str(path))
    try:
        for name, source_name in SOURCE_TABLE_NAMES.items():
            df = _frame(name, tables.get(name, []))
            con.register("_t", df)
            con.execute(f"CREATE TABLE {source_name} AS SELECT {_typed_select(name)} FROM _t")
            con.unregister("_t")
    finally:
        con.close()
    return path


# --- row builders: every field defaults to a valid, eligible value -------------------------

def co_filing(fileid: str, filingdate: str | None = "2001-02-03T00:00:00.000", master: str | None = None,
              filingtype: str = "ucc", terminationflag: str = "false") -> dict:
    return {"transactionid": f"T-{fileid}", "masterdocumentid": master or f"M-{fileid}",
            "transactiontype": "UCC1", "filingtype": filingtype, "documenttype": "UCC1",
            "filingdate": filingdate, "continuation": "false", "terminationflag": terminationflag,
            "fileid": fileid}


def co_debtor(fileid: str, name: str | None, address: str | None = "1 MAIN ST", city: str | None = "DENVER",
              state: str | None = "CO", zipcode: str | None = "80202", debtorid: str | None = None) -> dict:
    return {"debtorid": debtorid or f"D-{fileid}-{name}", "organizationname": name, "address1": address,
            "city": city, "state": state, "zipcode": zipcode, "fileid": fileid, "actiontype": "ADD",
            "recordstatus": "A", "efsuniqueid": None}


def co_sp(fileid: str, name: str | None) -> dict:
    return {"spid": f"S-{fileid}-{name}", "organizationname": name, "address1": "9 BANK ST", "city": "DENVER",
            "state": "CO", "zipcode": "80202", "fileid": fileid, "actiontype": "ADD", "recordstatus": "A",
            "assignor": None}


def ct_row(lien_id: str | None, name: str | None, address: str | None = "1 MAIN ST", city: str | None = "HARTFORD",
           state: str | None = "CT", zipcode: str | None = "06103", lender: str | None = "FIRST BANK",
           dt_accept: str | None = "2001-02-03T00:00:00.000", status: str = "Active",
           ftype: str = "ORIG FIN STMT") -> dict:
    return {"id_lien_flng_nbr": lien_id, "lien_status": status, "cd_flng_type": ftype, "debtor_nm_bus": name,
            "debtor_ad_str1": address, "debtor_ad_city": city, "debtor_ad_state": state, "debtor_ad_zip": zipcode,
            "sec_party_nm_bus": lender, "sec_party_ad_str1": "9 BANK ST", "sec_party_ad_city": "HARTFORD",
            "sec_party_ad_state": "CT", "sec_party_ad_zip": "06103", "dt_lapse": None, "dt_accept": dt_accept}


def scope_row(region: str, fileid: str, borrower: str, lender: str | None, address: str = "1 MAIN ST",
              city: str = "DENVER", state: str = "CO", zipcode: str = "80202", loan_year: str = "2001",
              route_a: bool = False, route_b: bool = True) -> dict:
    return {"region": region, "fileid": fileid, "loan_year": loan_year, "borrower": borrower,
            "borrower_address": address, "borrower_city": city, "borrower_state": state, "borrower_zip": zipcode,
            "lender": lender, "route_a": route_a, "route_b": route_b}


@pytest.fixture
def snapshot_factory(tmp_path):
    def _make(tables: dict[str, list[dict]], name: str = "snap") -> Path:
        return write_snapshot(tmp_path / name, tables)
    return _make
# --- synthetic Case rows (valid under contracts.Case, scored by the real vendored baseline) ----

def make_case_row(region: str, file_id: str, name: str, lenders=(), address: str = "1 MAIN ST",
                  city: str | None = None, state: str | None = None, zipcode: str | None = None,
                  dataset_version: str = "v1") -> dict:
    from datetime import date

    from ucc_ml import legacy
    from ucc_ml.contracts import make_borrower_key, make_case_id, make_group_id
    from ucc_ml.dataset import canonical_lender_set

    city = city or ("DENVER" if region == "CO" else "HARTFORD")
    state = state or region
    zipcode = zipcode or ("80202" if region == "CO" else "06103")
    lenders = canonical_lender_set(lenders)
    name_clean, suffix = legacy.normalize_name(name)
    bk = make_borrower_key(name, address, city, state, zipcode)
    return {
        "case_id": make_case_id(region, file_id, bk), "dataset_version": dataset_version, "region": region,
        "file_id": file_id, "lineage_id": None, "borrower_key": bk, "borrower_name_raw": name,
        "borrower_name_clean": name_clean, "borrower_suffix": suffix,
        "lender_names_raw": lenders,
        "lender_names_clean": sorted({c for c in (legacy.normalize_name(l)[0] for l in lenders) if c}),
        "earliest_observed_date": date(2001, 2, 3), "latest_observed_date": date(2001, 2, 3),
        "borrower_city": city, "borrower_state": state, "borrower_zip": zipcode,
        "source_row_count": 1, "source_filing_type": "ucc" if region == "CO" else "ORIG FIN STMT",
        "source_status": "false" if region == "CO" else "Active",
        "baseline_qualifies": legacy.baseline_qualifies(name, lenders),
        "baseline_route": legacy.baseline_route(name, lenders),
        "group_id": make_group_id(region, name_clean, name),
    }


def write_candidates_fixture(path: Path, rows: list[dict]) -> pd.DataFrame:
    from ucc_ml.contracts import CASE_COLUMNS
    from ucc_ml.dataset import validate_cases, write_candidates

    df = pd.DataFrame(rows, columns=list(CASE_COLUMNS)).sort_values("case_id").reset_index(drop=True)
    for c in ("lineage_id", "borrower_name_clean", "borrower_suffix", "source_filing_type", "source_status",
              "borrower_city", "borrower_state", "borrower_zip"):
        df[c] = pd.Series([v if isinstance(v, str) else None for v in df[c]], index=df.index, dtype=object)
    validate_cases(df)
    write_candidates(df, path)
    return df


def balanced_case_rows(per_stratum: int) -> list[dict]:
    """per_stratum accepted + per_stratum rejected cases in each region; every case its own group."""
    rows = []
    for region in ("CO", "CT"):
        for i in range(per_stratum):
            rows.append(make_case_row(region, f"{region}A{i:04d}", f"ACME EXCAVATION {i} LLC", ["FIRST BANK"]))
            rows.append(make_case_row(region, f"{region}R{i:04d}", f"SMITH LAW OFFICE {i} PC", ["FIRST BANK"]))
    return rows


# --- synthetic labelling rounds (Tasks 15-21): fake blind passes and a tmp repository with a queued pilot ---
from collections.abc import Iterable  # noqa: E402

FLIPPED_ANSWER = {"label": "INSUFFICIENT_EVIDENCE", "reason_code": "I_AMBIGUOUS_WORD",
                  "reason": "reads as a generic firm name"}


def fake_answer(borrower_name: str) -> dict:
    """A deterministic stand-in for one blind labeller decision."""
    name = borrower_name.upper()
    if "EXCAVATION" in name:
        return {"label": "RELEVANT", "reason_code": "R_BORROWER_EQUIPMENT_WORD", "reason": "EXCAVATION in the borrower name"}
    if "LAW OFFICE" in name:
        return {"label": "NOT_RELEVANT", "reason_code": "N_OTHER_INDUSTRY_EXPLICIT",
                "reason": "LAW OFFICE states a different activity"}
    return {"label": "INSUFFICIENT_EVIDENCE", "reason_code": "I_GENERIC_NAME", "reason": "the name could be anything"}


def fake_structured_output(chunk: pd.DataFrame, flips: Iterable[str] = ()) -> dict:
    """{"rows": [...]} for one queue chunk; a row whose queue case_id or borrower name is in `flips`
    gets FLIPPED_ANSWER instead of fake_answer."""
    flips = set(flips)
    rows = []
    for case_id, name in zip(chunk.case_id, chunk.borrower_name_raw):
        answer = FLIPPED_ANSWER if (case_id in flips or name in flips) else fake_answer(name)
        rows.append({"case_id": case_id, **answer})
    return {"rows": rows}


def labelling_repo(tmp_path: Path, per_stratum: int = 10, chunk_size: int = 15, repeat_fraction: float = 0.10,
                   audit_per_split_stratum: int = 2, extra_per_stratum: int = 20) -> Path:
    """A tmp repository with config, candidates, a pilot, frozen splits, FROZEN specs and the queued pilot round.

    Runs the real make-pilot, freeze-splits and label-blind commands and returns the config path."""
    import shutil

    from test_config import MINIMAL
    from ucc_ml.cli import main

    cfg = tmp_path / "ml/configs/v1.yaml"
    cfg.parent.mkdir(parents=True, exist_ok=True)
    cfg.write_text(MINIMAL.replace("pilot_per_stratum: 50", f"pilot_per_stratum: {per_stratum}")
                   .replace("chunk_size: 200", f"chunk_size: {chunk_size}")
                   .replace("repeat_fraction: 0.10", f"repeat_fraction: {repeat_fraction}")
                   .replace("founder_audit_per_split_stratum: 10",
                            f"founder_audit_per_split_stratum: {audit_per_split_stratum}"))
    cand = tmp_path / "ml/data/candidates/v1/candidates.parquet"
    cand.parent.mkdir(parents=True, exist_ok=True)
    write_candidates_fixture(cand, balanced_case_rows(per_stratum + extra_per_stratum))
    specs = tmp_path / "ml/specs"
    specs.mkdir(parents=True, exist_ok=True)
    real_specs = Path(__file__).resolve().parents[1] / "specs"
    policy = (real_specs / "label_policy_v1.md").read_text(encoding="utf-8")
    policy = "\n".join("status: FROZEN" if line.startswith("status:") else line for line in policy.splitlines())
    (specs / "label_policy_v1.md").write_text(policy + "\n", encoding="utf-8")
    shutil.copy(real_specs / "labeller_prompt_v1.md", specs / "labeller_prompt_v1.md")
    for command in (["make-pilot"], ["freeze-splits"], ["label-blind", "--round", "pilot_v1"]):
        assert main([command[0], "--config", str(cfg), *command[1:]]) == 0, command
    return cfg


def run_blind_passes(cfg: Path, round_name: str, flips_in_pass_b: Iterable[str] = ()) -> None:
    """Both fake passes for every chunk of a round, through the real write-raw-labels command."""
    import json

    from ucc_ml.cli import main
    from ucc_ml.config import load_config
    from ucc_ml.labeling import queue_parts, read_key, read_queue, round_paths

    rp = round_paths(load_config(cfg), round_name)
    flips_in_pass_b = list(flips_in_pass_b)
    for part in queue_parts(read_key(rp.key)):
        chunk = read_queue(rp.queue_part(part))
        for letter, flips in (("a", ()), ("b", flips_in_pass_b)):
            path = rp.structured_output(letter, part)
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(json.dumps(fake_structured_output(chunk, flips)), encoding="utf-8")
            assert main(["write-raw-labels", "--config", str(cfg), "--round", round_name, "--pass", letter,
                         "--part", str(part), "--structured", str(path)]) == 0


def imported_pilot_repo(tmp_path: Path, n_disagreements: int = 3, **repo_kwargs) -> tuple[Path, list[str]]:
    """labelling_repo + both fake passes (pass B flips the first n original cases of part 1) + import-labels.

    Returns the config path and the case_ids on which the two passes disagree."""
    from ucc_ml.cli import main
    from ucc_ml.config import load_config
    from ucc_ml.labeling import read_key, round_paths

    cfg = labelling_repo(tmp_path, **repo_kwargs)
    key = read_key(round_paths(load_config(cfg), "pilot_v1").key)
    flips = key[(key.part == 1) & ~key.is_repeat].queue_case_id.tolist()[:n_disagreements]
    run_blind_passes(cfg, "pilot_v1", flips_in_pass_b=flips)
    assert main(["import-labels", "--config", str(cfg), "--round", "pilot_v1"]) == 0
    return cfg, flips


def fill_founder_workbook(path: Path, answers: dict) -> None:
    """Write founder_label / founder_note into the REVIEW THESE sheet for the case_ids in `answers`."""
    from openpyxl import load_workbook

    from ucc_ml.labeling import WORKBOOK_COLUMNS, WORKBOOK_SHEET

    workbook = load_workbook(path)
    sheet = workbook[WORKBOOK_SHEET]
    label_column = WORKBOOK_COLUMNS.index("founder_label") + 1
    note_column = WORKBOOK_COLUMNS.index("founder_note") + 1
    for row in range(2, sheet.max_row + 1):
        case_id = sheet.cell(row=row, column=WORKBOOK_COLUMNS.index("case_id") + 1).value
        if case_id in answers:
            label, note = answers[case_id]
            sheet.cell(row=row, column=label_column, value=label)
            sheet.cell(row=row, column=note_column, value=note or None)
    workbook.save(path)


def founder_answers(cfg: Path, round_name: str, decide_all: bool = True) -> dict:
    """Answers for a round's issued review: every disagreement decided (all but the first when not
    decide_all), the first audit row confirmed and the second overturned."""
    import json

    from ucc_ml.config import load_config
    from ucc_ml.labeling import read_pass_file, round_paths

    rp = round_paths(load_config(cfg), round_name)
    manifest = json.loads(rp.review_manifest.read_text())
    pass_a = read_pass_file(rp.pass_file("a"))
    agreed = pass_a[~pass_a.is_repeat].set_index("case_id").label
    decided = manifest["disagreements"] if decide_all else manifest["disagreements"][1:]
    answers = {case_id: ("NOT_RELEVANT", "founder: the name states a different activity") for case_id in decided}
    audit = manifest["audit"]
    if audit:
        answers[audit[0]] = (agreed[audit[0]], "")
    if len(audit) > 1:
        other = "INSUFFICIENT_EVIDENCE" if agreed[audit[1]] != "INSUFFICIENT_EVIDENCE" else "RELEVANT"
        answers[audit[1]] = (other, "founder: the name alone does not settle it")
    return answers


def reviewed_pilot_repo(tmp_path: Path, decide_all: bool = True, **kwargs) -> Path:
    """imported_pilot_repo + review-workbook + founder answers + import-founder-review; returns the config path."""
    from ucc_ml.cli import main
    from ucc_ml.config import load_config
    from ucc_ml.labeling import round_paths

    cfg, _ = imported_pilot_repo(tmp_path, **kwargs)
    assert main(["review-workbook", "--config", str(cfg), "--round", "pilot_v1"]) == 0
    fill_founder_workbook(round_paths(load_config(cfg), "pilot_v1").workbook, founder_answers(cfg, "pilot_v1", decide_all))
    assert main(["import-founder-review", "--config", str(cfg), "--round", "pilot_v1"]) == (0 if decide_all else 1)
    return cfg
